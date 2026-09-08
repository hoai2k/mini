"""Build Hopper_3D_Design.docx from source/design.md, in the same layout as the
2D design package (hopper/design/source/build_document.py). Then convert to
PDF with LibreOffice if it is installed:
    python3 build_document.py && soffice --headless --convert-to pdf ../Hopper_3D_Design.docx --outdir ..
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'source/design.md'
text = source.read_text()
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.64); sec.bottom_margin = Inches(.64)
sec.left_margin = Inches(.68); sec.right_margin = Inches(.68)
sec.header_distance = Inches(.27); sec.footer_distance = Inches(.27)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet', 'List Number']:
    st = doc.styles[name]; st.font.name = 'Arial'; st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(7)
st = doc.styles['Normal']; st.font.size = Pt(11); st.paragraph_format.line_spacing = 1.06
for name, size in [('Title', 29), ('Subtitle', 15), ('Heading 1', 21), ('Heading 2', 13), ('Heading 3', 12)]:
    st = doc.styles[name]; st.font.size = Pt(size); st.font.bold = name not in ['Subtitle']
    st.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(9)
for name in ['List Bullet', 'List Number']:
    doc.styles[name].font.size = Pt(10.5); doc.styles[name].paragraph_format.space_after = Pt(3)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.color.rgb = RGBColor.from_string('505A54')
sec.different_first_page_header_footer = True
h = sec.header.paragraphs[0]; h.text = 'HOPPER THE GRASSHOPPER     /     3D EDITION     /     GAME DESIGN AND PLAN'
h.runs[0].font.size = Pt(8); h.runs[0].font.color.rgb = RGBColor(0, 0, 0)
f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = f.add_run('DESIGN REVIEW     •     '); r.font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); f._p.append(fld)

def inline(p, s, size=None):
    """Add runs for **bold** and `code` spans."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', s):
        if not part: continue
        if part.startswith('**'): run = p.add_run(part[2:-2]); run.bold = True
        elif part.startswith('`'): run = p.add_run(part[1:-1]); run.font.name = 'Consolas'
        else: run = p.add_run(part)
        if size: run.font.size = Pt(size)

def table(rows):
    n = len(rows[0])
    t = doc.add_table(rows=1, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    widths = {2: [2.0, 5.14], 3: [2.15, 2.4, 2.59], 4: [1.7, 1.65, 1.7, 2.09], 5: [1.5, 1.4, 1.4, 1.4, 1.44], 6: [1.3, 1.15, 1.15, 1.15, 1.15, 1.24]}.get(n, [7.14 / n] * n)
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, s) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '283B50' if i == 0 else ('F2F4F7' if i % 2 == 0 else 'FFFFFF')); tcPr.append(sh)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            tcPr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, v in [('top', 80), ('bottom', 80), ('left', 100), ('right', 100)]:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcPr.append(margins)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.02
            inline(p, s, size=9.5 if n > 3 else 10)
            for run in p.runs:
                run.bold = run.bold or i == 0; run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '111111')
        pr = t.rows[i]._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        if i == 0: pr.append(OxmlElement('w:tblHeader'))
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = Pt(2); p.add_run(' ').font.size = Pt(2)

pages = text.split('<!-- page -->')
for pageidx, page in enumerate(pages):
    if pageidx: doc.add_page_break()
    lines = page.strip().splitlines(); i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [s.strip() for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[-: ]+', s) for s in cells): rows.append(cells)
                i += 1
            table(rows); continue
        m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if m:
            path = (source.parent / m.group(2)).resolve()
            width = 7.14
            if 'canonical-v1' in path.name or 'logo' in path.name: width = 6.4
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            shape = p.add_run().add_picture(str(path), width=Inches(width))
            shape._inline.docPr.set('descr', m.group(1))
            p.paragraph_format.keep_with_next = True
            cap = doc.add_paragraph(m.group(1), 'Caption'); cap.paragraph_format.space_after = Pt(9)
        elif line.startswith('# '): doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '): doc.add_paragraph(line[3:], 'Subtitle' if pageidx == 0 else 'Heading 1')
        elif line.startswith('### '): doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('- '): inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line): inline(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        else:
            p = doc.add_paragraph(); inline(p, line)
        i += 1

doc.core_properties.title = 'Hopper the Grasshopper — 3D Edition — Game Design and Plan'
doc.core_properties.subject = '3D reinterpretation: design, production plan and asset requests'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Hopper, game design, 3D, three.js, Xbox, anime'
for tree in [doc.styles.element, doc.element]:
    for border in list(tree.iter(qn('w:pBdr'))): border.getparent().remove(border)
out = ROOT / 'Hopper_3D_Design.docx'; doc.save(out)
print(f'Created {out}; {len(pages)} designed sections; {len(text.split())} words')
